"""python3 -m unittest discover engine/tests   (from the repo root)

The checker is the release gate, so most tests break a valid schedule on purpose
and assert the checker catches it.
"""
import copy
import sys
import unittest
from pathlib import Path

ENGINE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ENGINE))

from adapters import factory as fa  # noqa: E402
from core import greedy  # noqa: E402
from core.check import check as check_schedule  # noqa: E402
from model import load_config, setup_minutes  # noqa: E402

COFFEE = (ENGINE / "configs/coffee_roastery.yaml", ENGINE / "samples/coffee_orders.csv", "2026-10-18T08:00")
NAHDA = (ENGINE / "configs/nahda_riyadh.yaml", ENGINE / "samples/nahda_orders.csv", "2026-10-18T06:00")


def plan(case):
    cfg, orders, start = case
    prob = fa.build(cfg, orders, start)
    return prob, greedy.solve(prob)[0]


class TestModel(unittest.TestCase):
    def test_toll_100kg_is_six_full_and_two_of_eight(self):
        prob, _ = plan(COFFEE)
        sizes = sorted(t.qty for t in prob.meta["factory"].tasks.values() if "toll" in t.id)
        self.assertEqual(sizes, [8, 8, 14, 14, 14, 14, 14, 14])

    def test_grinder_setup_is_asymmetric(self):
        cfg = load_config(COFFEE[0])
        g = cfg["resources"]["grinder_1"]
        self.assertEqual(setup_minutes(g, {"flavor": "cardamom"}, {"flavor": "plain"}), 30)
        self.assertEqual(setup_minutes(g, {"flavor": "plain"}, {"flavor": "cardamom"}), 0)
        self.assertEqual(setup_minutes(g, {"flavor": "spiced"}, {"flavor": "cardamom"}), 0)

    def test_filler_setup_adds_size_change(self):
        cfg = load_config(NAHDA[0])
        f = cfg["resources"]["line_pet"]
        self.assertEqual(setup_minutes(f, {"family": "orange", "size_ml": 1500}, {"family": "water", "size_ml": 500}), 85)
        self.assertEqual(setup_minutes(f, {"family": "cola", "size_ml": 500}, {"family": "lemon", "size_ml": 500}), 90)


class TestChecker(unittest.TestCase):
    def test_rule_plans_pass(self):
        for case in (COFFEE, NAHDA):
            prob, sched = plan(case)
            self.assertEqual(check_schedule(prob, sched), [])

    def _broken(self, mutate):
        prob, sched = plan(COFFEE)
        bad = copy.deepcopy(sched)
        mutate(prob, bad)
        return check_schedule(prob, bad)

    def test_catches_overlap(self):
        def f(prob, s):
            roast = sorted((v["start"], k) for k, v in s.items() if v["res"][0] == "roaster_1")
            a, b = roast[0][1], roast[1][1]
            d = s[b]["end"] - s[b]["start"]
            s[b]["start"] = s[a]["start"] + 10
            s[b]["end"] = s[b]["start"] + d
        self.assertTrue(any("overlaps" in v for v in self._broken(f)))

    def test_catches_rest_time(self):
        def f(prob, s):
            blend = next(k for k in s if k.endswith("-blend"))
            pred = next(a for a, b, lo, hi in prob.preds if b == blend)
            d = s[blend]["end"] - s[blend]["start"]
            s[blend]["start"] = s[pred]["end"] + 5
            s[blend]["end"] = s[blend]["start"] + d
        self.assertTrue(any("minimum is 60" in v for v in self._broken(f)))

    def test_catches_mixed_owners(self):
        def f(prob, s):
            toll = next(t for t in prob.meta["factory"].tasks.values() if "toll" in t.id)
            toll.loads.append(("O01", "green_brazil", 0.5))
        self.assertTrue(any("mixes material" in v for v in self._broken(f)))

    def test_catches_night_work(self):
        def f(prob, s):
            k = next(k for k, v in s.items() if v["res"][0] == "roaster_1")
            d = s[k]["end"] - s[k]["start"]
            s[k]["start"], s[k]["end"] = 1140, 1140 + d       # 03:00, the factory is closed
        self.assertTrue(any("outside working hours" in v for v in self._broken(f)))


if __name__ == "__main__":
    unittest.main()


class TestExcelAndBacktest(unittest.TestCase):
    XLSX = ENGINE / "samples/factory_log_example.xlsx"

    def test_import_reads_every_order(self):
        from openpyxl import load_workbook
        import import_excel as ix
        cfg = load_config(COFFEE[0])
        orders, problems = ix.read_orders(load_workbook(self.XLSX, data_only=True), cfg)
        self.assertEqual(len(orders), 35)
        self.assertEqual(sum(o["rejected"] for o in orders), 1)
        self.assertEqual(problems, [])

    def test_backtest_finishes_orders_without_peeking(self):
        from datetime import date
        from openpyxl import load_workbook
        import import_excel as ix
        from backtest import run
        cfg = load_config(COFFEE[0])
        wb = load_workbook(self.XLSX, data_only=True)
        rows = [o for o in ix.read_orders(wb, cfg)[0] if not o["rejected"]]
        done, batches = run(cfg, rows, ix.read_events(wb), date(2026, 10, 3), date(2026, 10, 16), "rules", 1)
        self.assertGreater(len(done), 25)
        for oid, t in done.items():                       # nothing finishes before it was ordered
            r = next(r for r in rows if r["order_id"] == oid)
            self.assertGreater(t.isoformat(), r["release"])
        self.assertTrue(all(8 - 1e-6 <= b["kg"] <= 14 + 1e-6 for b in batches))
